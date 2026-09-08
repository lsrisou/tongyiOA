import ssh
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

cursor = ssh.conn.cursor()

level_map = {'30': '超级事业部', '60': '校区合伙人'}


def fetch_schools(limit=None):
    sql = """
        SELECT id, name, level, join_at, deleted_at
        FROM deshengoa.de_schools
        WHERE company_id = 0 AND deleted_at IS NULL
        ORDER BY id ASC
    """
    if limit:
        sql += " LIMIT %s"
        cursor.execute(sql, (limit,))
    else:
        cursor.execute(sql)
    columns = [desc[0] for desc in cursor.description]
    rows = cursor.fetchall()

    school_ids = [dict(zip(columns, rs))['id'] for rs in rows]

    phone_map = {}
    if school_ids:
        placeholders = ','.join(['%s'] * len(school_ids))
        cursor.execute(f"""
            SELECT t.school_id, t.username FROM (
                SELECT school_id, username, MIN(id) as min_id
                FROM deshengoa.de_admin_users
                WHERE school_id IN ({placeholders}) AND deleted_at IS NULL
                GROUP BY school_id
            ) t
        """, school_ids)
        for school_id, username in cursor.fetchall():
            phone_map[school_id] = username

    schools = []
    for rs in rows:
        s = dict(zip(columns, rs))
        level_val = str(s['level'])
        level_display = level_map.get(level_val, level_val)

        join_at = s.get('join_at')
        if isinstance(join_at, datetime):
            join_at_str = join_at.strftime('%Y-%m-%d %H:%M:%S')
        elif isinstance(join_at, (int, float)) and join_at > 0:
            join_at_str = datetime.fromtimestamp(join_at).strftime('%Y-%m-%d %H:%M:%S')
        else:
            join_at_str = join_at if join_at else ""

        phone = phone_map.get(s['id'], "")
        if phone and len(phone) >= 7:
            phone = phone[:3] + "****" + phone[-4:]

        schools.append({
            '校区ID': s['id'],
            '校区名称': s['name'],
            '联系电话': phone,
            '校区级别': level_display,
            '加入时间': join_at_str,
        })
    return schools


def export_to_excel(schools, limit=None):
    label = f"前{limit}条" if limit else "全量"
    filename = f"校区基础信息表({label})-{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "校区基础信息"

    headers = ['校区ID', '校区名称', '联系电话', '校区级别', '加入时间']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, s in enumerate(schools, 2):
        for col_idx, key in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=s[key])
            cell.alignment = Alignment(horizontal='center')

    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 20

    wb.save(filename)
    print(f"导出成功: {filename}")
    print(f"共 {len(schools)} 条记录")


if __name__ == '__main__':
    schools = fetch_schools()
    export_to_excel(schools)
