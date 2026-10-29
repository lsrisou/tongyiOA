import ssh
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

def extract_income_source(msg):
    if not msg:
        return ''
    text = re.sub(r'^\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}', '', msg).strip()
    if '，学生：' in text:
        text = text.split('，学生：')[0]
    if '，代理：' in text:
        text = text.split('，代理：')[0]
    return text

def fmt_time(dt):
    return str(dt).split('.')[0] if dt else ''

def export_month(year, month, cursor):
    month_str = f'{year}-{month:02d}'
    start = f'{month_str}-01 00:00:00'
    end = f'{year}-{month+1:02d}-01 00:00:00' if month < 12 else f'{year+1}-01-01 00:00:00'

    top_sql = f"""
        SELECT top.id, top.amount, top.created_at, top.hour_lesson_id, top.student_id,
               s.name AS student_name
        FROM `deshengoa`.`de_top_point_logs` top
        JOIN `deshengoa`.`de_student_hour_lessons` shl ON top.hour_lesson_id = shl.id
        JOIN `deshengoa`.`de_students` s ON shl.student_id = s.id
        WHERE top.type_id = 4
          AND top.created_at >= '{start}'
          AND top.created_at < '{end}'
        ORDER BY top.created_at ASC
    """
    cursor.execute(top_sql)
    columns = [desc[0] for desc in cursor.description]
    top_rows = [dict(zip(columns, row)) for row in cursor.fetchall()]

    if not top_rows:
        print(f'{month_str} 无数据')
        return

    hour_lesson_ids = list(set(row['hour_lesson_id'] for row in top_rows))
    ids_str = ','.join(str(i) for i in hour_lesson_ids)

    user_sql = f"""
        SELECT u.amount, u.created_at, u.hour_lesson_id, u.msg,
               s.name AS student_name, shl.student_id
        FROM `deshengoa`.`de_user_point_logs` u
        JOIN `deshengoa`.`de_student_hour_lessons` shl ON u.hour_lesson_id = shl.id
        JOIN `deshengoa`.`de_students` s ON shl.student_id = s.id
        WHERE u.hour_lesson_id IN ({ids_str})
        ORDER BY u.created_at ASC
    """
    cursor.execute(user_sql)
    columns = [desc[0] for desc in cursor.description]
    user_rows = [dict(zip(columns, row)) for row in cursor.fetchall()]

    school_sql = f"""
        SELECT sc.amount, sc.created_at, sc.hour_lesson_id, sc.msg,
               s.name AS student_name, shl.student_id
        FROM `deshengoa`.`de_school_point_logs` sc
        JOIN `deshengoa`.`de_student_hour_lessons` shl ON sc.hour_lesson_id = shl.id
        JOIN `deshengoa`.`de_students` s ON shl.student_id = s.id
        WHERE sc.hour_lesson_id IN ({ids_str})
        ORDER BY sc.created_at ASC
    """
    cursor.execute(school_sql)
    columns = [desc[0] for desc in cursor.description]
    school_rows = [dict(zip(columns, row)) for row in cursor.fetchall()]

    rows = []

    for r in top_rows:
        rows.append({
            '学生姓名': r['student_name'],
            '学生ID': r['student_id'],
            '课程ID': r['hour_lesson_id'],
            '收入来源': '总部教练课时抽成',
            '描述': '',
            '金额(元)': round(r['amount'] / 100, 2),
            '时间': fmt_time(r['created_at']),
        })

    for r in user_rows:
        rows.append({
            '学生姓名': r['student_name'],
            '学生ID': r['student_id'],
            '课程ID': r['hour_lesson_id'],
            '收入来源': extract_income_source(r['msg']),
            '描述': r['msg'],
            '金额(元)': round(r['amount'] / 100, 2),
            '时间': fmt_time(r['created_at']),
        })

    for r in school_rows:
        rows.append({
            '学生姓名': r['student_name'],
            '学生ID': r['student_id'],
            '课程ID': r['hour_lesson_id'],
            '收入来源': extract_income_source(r['msg']),
            '描述': r['msg'],
            '金额(元)': round(r['amount'] / 100, 2),
            '时间': fmt_time(r['created_at']),
        })

    order_map = {}
    for idx, r in enumerate(top_rows):
        key = (r['student_name'], r['hour_lesson_id'])
        if key not in order_map:
            order_map[key] = idx

    rows.sort(key=lambda x: (
        order_map.get((x['学生姓名'], x['课程ID']), 999999),
        1 if x['收入来源'] == '总部教练课时抽成' else 0,
        x['时间'],
    ))

    wb = Workbook()
    ws = wb.active
    ws.title = '课程分成报表'

    headers = ['学生姓名', '学生ID', '课程ID', '收入来源', '描述', '金额(元)', '时间']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.font = Font(bold=True)

    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = row_data[header]
            if header in ('金额(元)',):
                cell.alignment = Alignment(horizontal='right')
            elif header in ('描述', '收入来源'):
                cell.alignment = Alignment(horizontal='left', wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal='center')

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 20

    filename = f'课程分成报表-{month_str}.xlsx'
    wb.save(filename)
    print(f'{filename} 导出成功（top: {len(top_rows)}, user: {len(user_rows)}, school: {len(school_rows)}，共 {len(rows)} 条）')


cursor = ssh.conn.cursor()

for month in [9]:
    export_month(2026, month, cursor)

cursor.close()
ssh.conn.close()
print('全部完成')
